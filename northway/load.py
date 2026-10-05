"""Read the six export files exactly as Sam sent them.

Every loader is read-only and fixes that file's specific mess in code,
so nobody has to hand-edit anything.
"""
import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

SKU_BASE = re.compile(r"^([A-Z]+-\d+)")
OLD_NW_CODE = re.compile(r"^NW-([A-Z]+)(\d+)$")  # e.g. NW-HTW152 -> HTW-152


def sku_base(sku):
    """'WCL-101-6' -> 'WCL-101'. Returns None if it doesn't look like a SKU."""
    m = SKU_BASE.match(str(sku or "").strip().upper())
    return m.group(1) if m else None


def money(text):
    """'$1,220.99' -> 1220.99 ; '' -> None"""
    t = str(text or "").replace("$", "").replace(",", "").strip()
    return float(t) if t else None


def number(text):
    """'2,400.00' -> 2400.0 ; '' -> None"""
    t = str(text or "").replace(",", "").strip()
    return float(t) if t else None


# --------------------------------------------------------------- pricing.xlsx

@dataclass
class Sheet:
    products: list = field(default_factory=list)     # dicts, one per row
    prices: dict = field(default_factory=dict)       # base sku -> list of price rows
    shipments: list = field(default_factory=list)
    categories: list = field(default_factory=list)   # (name, multipliers or None)
    category_rules: list = field(default_factory=list)  # free-text rules under the table
    category_rows: list = field(default_factory=list)   # every named row, in order: (row number, name, multipliers)
    prices_generated: str = ""


def load_sheet(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    s = Sheet()

    # Products: row 1 is a merged title ("DO NOT SORT"), row 2 is the header.
    ws = wb["Products"]
    for row in ws.iter_rows(min_row=3):
        sku = row[0].value
        if not sku:
            continue
        fill = row[0].fill
        highlighted = bool(fill and fill.fill_type and fill.fgColor.rgb not in (None, "00000000"))
        s.products.append({
            "sku": str(sku).strip().upper(),
            "old_books_sku": (str(row[1].value).strip().upper() if row[1].value else None),
            "category": row[2].value,            # kept raw: spacing/case matters, see checks
            "pcs_per_box": row[3].value,
            "name": row[4].value,
            "notes": row[5].value,
            "highlighted": highlighted,          # Dana/Sam's pink rows
        })

    # Prices (auto): written nightly by the price program. Row 1 = timestamp, row 2 = header.
    ws = wb["Prices (auto)"]
    s.prices_generated = str(ws.cell(1, 1).value or "")
    for r in ws.iter_rows(min_row=3, values_only=True):
        if not r[0]:
            continue
        s.prices.setdefault(str(r[0]).strip().upper(), []).append({
            "status": r[1], "piece": r[3], "dozen": r[4], "box": r[5], "box3plus": r[6],
        })  # cost (r[2]) is deliberately not read: Marco and Leo shouldn't see costs.

    # Shipments: header in row 1. Dates are typed by hand in many formats, so keep them as text.
    ws = wb["Shipments"]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        s.shipments.append({
            "sku": str(r[0]).strip().upper(), "arrived": r[1], "pcs_per_case": r[5],
            "cases": r[7], "notes": r[8],
        })

    # Categories: a note row sits under the header, and business rules are typed below the table.
    # The price program takes the FIRST row whose name matches, so row order matters:
    # an empty row with a category's name (the "Bed Skirt" heading) hides the real row.
    ws = wb["Categories"]
    for n, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        name = r[0]
        if not name:
            continue
        mults = r[1:5]
        s.category_rows.append((n, str(name), None if all(m is None for m in mults) else mults))
        if all(m is None for m in mults):
            s.category_rules.append(str(name))   # a rule under the table, or an empty heading row
        else:
            s.categories.append((str(name), mults))
    return s


# ------------------------------------------------------------- books_item_list

def load_books(path):
    """Item Listing report. Has 3 title lines above the header and a TOTAL row at the end."""
    lines = Path(path).read_text(encoding="utf-8-sig").splitlines()
    as_of = next((l.strip('"') for l in lines[:4] if l.strip('"').startswith("As of")), "")
    start = next(i for i, l in enumerate(lines) if l.startswith("Item ID"))
    items = []
    for r in csv.DictReader(lines[start:]):
        if not r["Item ID"] or not r["Item ID"].isdigit():   # skips the TOTAL row
            continue
        items.append({
            "id": int(r["Item ID"]), "item": r["Item"].strip().upper(),
            "description": r["Description"], "type": r["Type"],
            "bundle_of": int(r["Bundle Of"]) if r["Bundle Of"] else None,
            "bundle_qty": number(r["Bundle Qty"]), "qty": number(r["Qty On Hand"]),
        })
    return items, as_of


# ---------------------------------------------------------------- web stores

def load_northway(path):
    """Store-admin export: only the first row of each product has Title/Status."""
    rows, title, status, handle = [], None, None, None
    with open(path, encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            if r["Handle"] != handle:
                handle, title, status = r["Handle"], r["Title"], r["Status"]
            pack = r["Option1 Value"]
            box = re.match(r"Box of (\d+)", pack)
            rows.append({
                "store": "Northway", "variant_id": int(r["Variant ID"]), "title": title,
                "status": status or "", "sku": r["Variant SKU"].strip().upper(),
                "pack": "box" if box else pack.lower(),            # piece / dozen / box
                "pack_size": int(box.group(1)) if box else (12 if pack == "Dozen" else 1),
                "price": money(r["Variant Price"]), "stock": number(r["Variant Inventory Qty"]),
            })
    return rows


def load_maple(path):
    """Exported from a different screen: other column names, $-strings, boxes only."""
    rows = []
    with open(path, encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            size = re.match(r"Box/(\d+)", r["Pack"] or "")
            rows.append({
                "store": "Maple", "variant_id": int(r["ID"]), "title": r["Product"],
                "status": "", "sku": r["SKU"].strip().upper(), "pack": "box",
                "pack_size": int(size.group(1)) if size else None,
                "price": money(r["Price CAD (Canada)"]), "stock": number(r["On hand"]),
            })
    return rows


def load_bridge(path):
    with open(path, encoding="utf-8-sig") as f:
        return [{"store": r["Store"], "variant_id": int(r["Store Variant ID"]),
                 "books_id": int(r["Books Item ID"]), "last_synced": r["Last Synced"]}
                for r in csv.DictReader(f)]


def load_all(exports_dir):
    d = Path(exports_dir)
    books, books_as_of = load_books(d / "books_item_list.csv")
    return {
        "sheet": load_sheet(d / "pricing.xlsx"),
        "books": books, "books_as_of": books_as_of,
        "web": load_northway(d / "northway_products_export.csv") + load_maple(d / "maple_export.csv"),
        "bridge": load_bridge(d / "bridge_mapping.csv"),
    }
